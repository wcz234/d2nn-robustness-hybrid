"""Build an editable Chinese Word manuscript from the reviewed Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PAPER_DIR = Path(__file__).resolve().parent
ROOT = PAPER_DIR.parent
SOURCE = PAPER_DIR / "MANUSCRIPT_ZH.md"
BIB = ROOT / "references.bib"
OUTPUT = PAPER_DIR / "MANUSCRIPT_ZH.docx"
EQUATION_DIR = PAPER_DIR / "docx_equations"

TITLE = "衍射深度神经网络鲁棒性与混合光电推理的协议约束数值评估：基于 MNIST 与 Fashion-MNIST 的研究"
FIGURES = {
    "figures/figure_1_method_overview.png": 15.5,
    "figures/figure_2_robustness.png": 15.8,
    "figures/figure_3_complexity.png": 15.8,
    "figures/figure_4_confusion.png": 15.5,
}


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=False, italic=False):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, east_asia="宋体", latin="Times New Roman", size=9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def parse_bib(path: Path) -> dict[str, dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    entries: dict[str, dict[str, str]] = {}
    for match in re.finditer(r"@\w+\{([^,]+),([\s\S]*?)\n\}", text):
        key, body = match.groups()
        fields: dict[str, str] = {}
        for field in ("author", "title", "journal", "booktitle", "year", "volume", "number", "pages", "doi"):
            field_match = re.search(rf"\b{field}\s*=\s*\{{([\s\S]*?)\}}", body)
            if field_match:
                fields[field] = re.sub(r"\s+", " ", field_match.group(1)).strip()
        entries[key.strip()] = fields
    return entries


def format_reference(index: int, entry: dict[str, str]) -> str:
    authors = entry.get("author", "").replace(" and ", ", ")
    title = entry.get("title", "")
    venue = entry.get("journal", entry.get("booktitle", ""))
    year = entry.get("year", "")
    volume = entry.get("volume", "")
    pages = entry.get("pages", "")
    details = ", ".join(x for x in (venue, volume, pages, year) if x)
    doi = entry.get("doi")
    result = f"[{index}] {authors}. {title}. {details}."
    if doi:
        result += f" DOI: {doi}."
    return result


def inline_plain(text: str, citation_numbers: dict[str, int]) -> str:
    def replace_cite(match):
        numbers = [str(citation_numbers[key.strip()]) for key in match.group(1).split(",") if key.strip() in citation_numbers]
        return "[" + ", ".join(numbers) + "]" if numbers else ""

    text = re.sub(r"\\cite\{([^}]+)\}", replace_cite, text)
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    text = text.replace("\\times", "×").replace("\\pi", "π").replace("\\times", "×")
    text = text.replace("\\&", "&")
    return text


def split_pipe_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def is_separator_row(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|"):
        return False
    cells = split_pipe_row(stripped)
    return bool(cells) and all(set(cell) <= {"-", ":", " "} and "-" in cell for cell in cells)


def add_table(doc: Document, caption_line: str, header: list[str], rows: list[list[str]], citation_numbers):
    """Render a Markdown pipe table as a real Word table, keeping the source caption text."""
    cap = doc.add_paragraph(style="Table Caption")
    add_inline_runs(cap, inline_plain(caption_line, citation_numbers), citation_numbers, size=9)

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.autofit = True
    for cell, value in zip(table.rows[0].cells, header):
        cell.text = ""
        add_inline_runs(cell.paragraphs[0], value, citation_numbers, size=9)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            add_inline_runs(cell.paragraphs[0], value, citation_numbers, size=9)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render_equation(index: int, source: str) -> Path:
    """Render a source LaTeX equation with Matplotlib's mathtext engine."""
    EQUATION_DIR.mkdir(exist_ok=True)
    path = EQUATION_DIR / f"equation_{index}.png"
    fig = plt.figure(figsize=(8.2, 0.7), dpi=300)
    fig.text(0.5, 0.5, f"${source.strip()}$", ha="center", va="center", fontsize=14)
    fig.patch.set_alpha(0)
    fig.savefig(path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    return path


def equation_width(path: Path) -> Inches:
    """Keep one math font size while fitting the longest equation in the page width."""
    width_px = Image.open(path).size[0]
    return Inches(min(width_px / 300.0, 6.05))


def add_inline_runs(paragraph, text: str, citation_numbers: dict[str, int], size=10.5):
    text = inline_plain(text, citation_numbers)
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, east_asia="等线", latin="Consolas", size=size)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for name, size, color, before, after in (("Heading 1", 15, "1F4E79", 14, 7), ("Heading 2", 12.5, "2F6690", 10, 5), ("Heading 3", 11.5, "2F6690", 8, 4)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    if "Equation" not in [s.name for s in doc.styles]:
        equation = doc.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = doc.styles["Equation"]
    equation.font.name = "Cambria Math"
    equation._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    equation.font.size = Pt(11)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(4)
    equation.paragraph_format.space_after = Pt(7)

    if "Figure Caption" not in [s.name for s in doc.styles]:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    if "Table Caption" not in [s.name for s in doc.styles]:
        table_caption = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_caption = doc.styles["Table Caption"]
    table_caption.font.name = "Times New Roman"
    table_caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    table_caption.font.size = Pt(9)
    table_caption.font.bold = False
    table_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_caption.paragraph_format.space_before = Pt(8)
    table_caption.paragraph_format.space_after = Pt(3)
    table_caption.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(TITLE)
    set_run_font(run, east_asia="黑体", latin="Times New Roman", size=18, bold=True)


def add_figure(doc: Document, image_path: Path, caption_text: str, citation_numbers):
    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.keep_with_next = True
    run = image_p.add_run()
    run.add_picture(str(image_path), width=Inches(6.15))
    cap = doc.add_paragraph(style="Figure Caption")
    add_inline_runs(cap, caption_text, citation_numbers, size=9)


def build():
    markdown = SOURCE.read_text(encoding="utf-8")
    bib_entries = parse_bib(BIB)
    cited_keys = []
    for key_group in re.findall(r"\\cite\{([^}]+)\}", markdown):
        for key in key_group.split(","):
            key = key.strip()
            if key and key not in cited_keys and key in bib_entries:
                cited_keys.append(key)
    citation_numbers = {key: i + 1 for i, key in enumerate(cited_keys)}

    doc = Document()
    configure_document(doc)
    add_title(doc)
    lines = markdown.splitlines()
    in_math = False
    math_lines: list[str] = []
    equation_index = 0
    index = 0
    seen_title = False
    while index < len(lines):
        line = lines[index]
        if line == "$$":
            if in_math:
                equation_index += 1
                equation = doc.add_paragraph(style="Equation")
                equation_path = render_equation(equation_index, " ".join(math_lines))
                equation.add_run().add_picture(str(equation_path), width=equation_width(equation_path))
                math_lines.clear()
            in_math = not in_math
            index += 1
            continue
        if in_math:
            math_lines.append(line)
            index += 1
            continue
        if not seen_title and line.startswith("# "):
            seen_title = True
            index += 1
            continue
        if line == "# 参考文献":
            break
        if not line.strip():
            index += 1
            continue
        if line.startswith("!["):
            image_match = re.fullmatch(r"!\[[^]]*\]\(([^)]+)\)", line)
            caption_line = lines[index + 2] if index + 2 < len(lines) and lines[index + 1] == "" else lines[index + 1]
            if image_match:
                image_path = PAPER_DIR / image_match.group(1)
                add_figure(doc, image_path, inline_plain(caption_line, citation_numbers), citation_numbers)
            index += 3 if index + 2 < len(lines) and lines[index + 1] == "" else 2
            continue
        heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            heading = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(heading, heading_match.group(2), citation_numbers, size={1: 15, 2: 12.5, 3: 11.5}[level])
            index += 1
            continue
        list_match = re.match(r"^(\d+)\.\s+(.+)$", line)
        if list_match:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            add_inline_runs(p, list_match.group(2), citation_numbers)
            index += 1
            continue
        caption_match = re.match(r"^\*\*图\s+\d+\s+\|", line)
        if caption_match:
            index += 1
            continue
        table_caption_match = re.match(r"^表\s+\d+\s*\|", line.strip())
        if table_caption_match:
            header_index = index + 1
            while header_index < len(lines) and not lines[header_index].strip():
                header_index += 1
            if header_index < len(lines) and lines[header_index].strip().startswith("|"):
                header = split_pipe_row(lines[header_index])
                separator_index = header_index + 1
                rows: list[list[str]] = []
                row_index = separator_index + 1
                while row_index < len(lines) and lines[row_index].strip().startswith("|"):
                    if not is_separator_row(lines[row_index]):
                        rows.append(split_pipe_row(lines[row_index]))
                    row_index += 1
                add_table(doc, line, header, rows, citation_numbers)
                index = row_index
                continue
        p = doc.add_paragraph()
        add_inline_runs(p, line, citation_numbers)
        index += 1

    ref_heading = doc.add_paragraph(style="Heading 1")
    add_inline_runs(ref_heading, "参考文献", citation_numbers, size=15)
    for i, key in enumerate(cited_keys, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(format_reference(i, bib_entries[key]))
        set_run_font(run, size=9)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "中文 SCI 研究论文主稿"
    doc.core_properties.comments = "Converted from the reviewed Chinese manuscript Markdown source."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
